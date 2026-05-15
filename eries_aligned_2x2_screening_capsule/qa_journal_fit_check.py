from pathlib import Path
from docx import Document

ROOT = Path(r"C:\Users\gjm31\OneDrive\Escritorio\articulos up\119 avo art\Entrega International Journal of Geotechnical Engineering 20260515")
MANUSCRIPT = ROOT / "01 Manuscript" / "Cyclic Lateral Pile Group - IJGE manuscript.docx"
COVER = ROOT / "05 Cover and declarations" / "Cover letter - IJGE.docx"
HIGHLIGHTS = ROOT / "05 Cover and declarations" / "Highlights - IJGE.docx"
FIG = ROOT / "02 Figures" / "Fig23_evidence_governed_workflow.png"
AUDIT = ROOT / "06 Journal instructions and audit" / "IJGE_editorial_fit_upgrade.md"

doc = Document(MANUSCRIPT)
text = "\n".join(p.text for p in doc.paragraphs)
title = doc.paragraphs[0].text.strip()
abstract = doc.paragraphs[5].text.strip()
assert title == "Evidence-Governed Pre-Design Screening of an ERIES-Aligned Laterally Loaded 2x2 Pile Group under Cyclic Degradation", title
assert len(abstract.split()) <= 250, len(abstract.split())
for required in [
    "evidence-governed pre-design screening workflow",
    "fixed ERIES-aligned 2x2 pile group",
    "FEM-anchored deterministic screening band",
    "not final design",
    "Table 1a. Fit-to-journal relevance map",
    "Figure 23. Evidence-governed pre-design workflow",
    "What the framework can support",
    "What the framework cannot support",
    "Why the workflow remains useful",
]:
    assert required in text, required

for forbidden in [
    "disruptive element",
    "absolute preliminary displacement band",
    "validated displacement prediction",
    "optimized topology",
    "confirmed by FEM",
    "commercial benchmark",
]:
    assert forbidden.lower() not in text.lower(), forbidden

cover_text = "\n".join(p.text for p in Document(COVER).paragraphs)
highlights_text = "\n".join(p.text for p in Document(HIGHLIGHTS).paragraphs)
assert title in cover_text
assert "field-calibrated displacement model" in cover_text
assert "preliminary displacement-screening band" in highlights_text
assert FIG.exists() and FIG.stat().st_size > 10000
assert AUDIT.exists() and "Recommended principal target" in AUDIT.read_text(encoding="utf-8")
print("JOURNAL_FIT_QA_PASSED")
